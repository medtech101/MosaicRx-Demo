"""
Verifies the Module 1 mandatory sanitization pass on a synthetic lecture
fixture containing every kind of leak the pipeline must catch: institution
name, course code, person name (with honorific), email address, an internal
URL, a repeated logo image, and a one-off image with blocklisted OCR text.
"""
import io
from pathlib import Path

import pytest
from PIL import Image, ImageDraw
from pptx import Presentation
from pptx.util import Inches

from app.services.sanitize.pipeline import run_ingestion_pipeline

FIXTURES_DIR = Path(__file__).parent / "fixtures"
FIXTURE_PATH = FIXTURES_DIR / "sample_lecture.pptx"

TERMS = {
    "institution": ["Springfield School of Medicine", "springfieldmed.edu"],
    "course_code": ["MED 501"],
    "person": ["Jane Smith"],
}


@pytest.fixture(scope="module")
def fixture_path() -> Path:
    FIXTURES_DIR.mkdir(parents=True, exist_ok=True)
    prs = Presentation()
    prs.core_properties.author = "Dr. Jane Smith"
    prs.core_properties.last_modified_by = "Jane Smith"
    prs.core_properties.comments = "Internal draft - do not distribute"

    master = prs.slide_masters[0]
    for layout in master.slide_layouts:
        for ph in layout.placeholders:
            try:
                if ph.placeholder_format.idx == 2 or "footer" in (ph.name or "").lower():
                    ph.text_frame.text = "Springfield School of Medicine - Confidential"
            except Exception:
                pass

    logo = Image.new("RGB", (100, 100), color=(10, 80, 200))
    ImageDraw.Draw(logo).text((10, 40), "LOGO", fill=(255, 255, 255))
    logo_buf = io.BytesIO()
    logo.save(logo_buf, format="PNG")
    logo_path = FIXTURES_DIR / "_logo.png"
    logo_path.write_bytes(logo_buf.getvalue())

    watermark = Image.new("RGB", (400, 200), color=(255, 255, 255))
    ImageDraw.Draw(watermark).text((10, 80), "Springfield School of Medicine", fill=(0, 0, 0))
    wm_buf = io.BytesIO()
    watermark.save(wm_buf, format="PNG")
    wm_path = FIXTURES_DIR / "_watermark.png"
    wm_path.write_bytes(wm_buf.getvalue())

    layout = prs.slide_layouts[1]

    slide1 = prs.slides.add_slide(layout)
    slide1.shapes.title.text = "Renal Physiology - Week 3"
    body = slide1.placeholders[1]
    body.text_frame.text = "Contact Dr. Jane Smith at jane.smith@springfieldmed.edu for questions."
    p = body.text_frame.add_paragraph()
    p.text = "See course MED 501 materials at https://portal.springfieldmed.edu/med501"
    slide1.shapes.add_picture(str(logo_path), Inches(0.2), Inches(0.2), height=Inches(0.5))
    slide1.notes_slide.notes_text_frame.text = "Remember to mention Professor Jane Smith reviewed this slide."

    slide2 = prs.slides.add_slide(layout)
    slide2.shapes.title.text = "The Renin-Angiotensin System"
    slide2.placeholders[1].text_frame.text = (
        "The renin-angiotensin system links renal, cardiovascular, and pharmacology physiology."
    )
    slide2.shapes.add_picture(str(logo_path), Inches(0.2), Inches(0.2), height=Inches(0.5))
    slide2.shapes.add_picture(str(wm_path), Inches(3), Inches(3), height=Inches(1.5))

    slide3 = prs.slides.add_slide(layout)
    slide3.shapes.title.text = "Hypertension pharmacology"
    slide3.placeholders[1].text_frame.text = (
        "ACE inhibitors reduce angiotensin II production, treating hypertension."
    )
    slide3.shapes.add_picture(str(logo_path), Inches(0.2), Inches(0.2), height=Inches(0.5))

    prs.save(FIXTURE_PATH)
    return FIXTURE_PATH


def test_unit_count(fixture_path):
    result = run_ingestion_pipeline(str(fixture_path), "pptx", TERMS)
    assert result.unit_count == 3


def test_body_text_is_redacted(fixture_path):
    result = run_ingestion_pipeline(str(fixture_path), "pptx", TERMS)
    assert "REDACTED" in result.sanitized_markdown
    assert "jane.smith@springfieldmed.edu" not in result.sanitized_markdown
    assert "MED 501" not in result.sanitized_markdown
    assert "springfieldmed.edu" not in result.sanitized_markdown
    assert "Renin-Angiotensin" in result.sanitized_markdown  # medical content preserved


def test_master_footer_never_reaches_output(fixture_path):
    result = run_ingestion_pipeline(str(fixture_path), "pptx", TERMS)
    assert "Confidential" not in result.sanitized_markdown
    assert "Springfield School of Medicine" not in result.sanitized_markdown


def test_metadata_is_flagged_for_scrubbing(fixture_path):
    result = run_ingestion_pipeline(str(fixture_path), "pptx", TERMS)
    fields = {r.location for r in result.metadata_redactions}
    assert "metadata:author" in fields
    assert "metadata:last_modified_by" in fields
    assert "metadata:comments" in fields


def test_repeated_logo_detected_and_watermark_flagged_for_review(fixture_path):
    result = run_ingestion_pipeline(str(fixture_path), "pptx", TERMS)
    logos = [d for d in result.image_decisions if d.is_logo]
    flagged = [d for d in result.image_decisions if d.needs_review]
    assert len(logos) == 3  # same logo bytes reused on all 3 slides
    assert len(flagged) == 1
    assert "Springfield School of Medicine" in flagged[0].matched_terms


def test_no_blocklist_leak_survives_in_final_markdown(fixture_path):
    """Defense in depth: even if every individual redaction worked, assert
    the whole markdown is clean end to end - this is what the confirm-time
    safety-net scan in the review router also checks before persisting."""
    from app.services.sanitize.blocklist import find_blocklist_matches

    result = run_ingestion_pipeline(str(fixture_path), "pptx", TERMS)
    leftover = find_blocklist_matches(result.sanitized_markdown, TERMS)
    assert leftover == []
