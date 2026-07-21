import hashlib
from docx import Document as DocxDocument

from .base import ParsedDocument, TextUnit, ExtractedImage

CORE_PROP_FIELDS = [
    "author", "category", "comments", "content_status", "identifier",
    "keywords", "language", "last_modified_by", "subject", "title",
    "version", "revision",
]


def parse_docx(path: str) -> ParsedDocument:
    d = DocxDocument(path)
    doc = ParsedDocument(doc_type="docx", unit_count=0)

    cp = d.core_properties
    for field_name in CORE_PROP_FIELDS:
        value = getattr(cp, field_name, None)
        if value:
            doc.core_properties[field_name] = str(value)
    company = getattr(cp, "company", None)
    if company:
        doc.core_properties["company"] = str(company)

    page_breaks = 1
    for i, para in enumerate(d.paragraphs):
        if para.text.strip():
            doc.text_units.append(TextUnit(location=f"paragraph {i + 1}", kind="body", text=para.text))
        if "\f" in (para.text or "") or any(run.font.name == "page-break" for run in para.runs):
            page_breaks += 1

    for t_idx, table in enumerate(d.tables, start=1):
        for row in table.rows:
            cells_text = [c.text for c in row.cells if c.text.strip()]
            if cells_text:
                doc.text_units.append(
                    TextUnit(location=f"table {t_idx}", kind="body", text="\n".join(cells_text))
                )

    for section_idx, section in enumerate(d.sections, start=1):
        for para in section.header.paragraphs:
            if para.text.strip():
                doc.text_units.append(
                    TextUnit(location=f"section {section_idx} header", kind="header", text=para.text)
                )
        for para in section.footer.paragraphs:
            if para.text.strip():
                doc.text_units.append(
                    TextUnit(location=f"section {section_idx} footer", kind="footer", text=para.text)
                )

    for rel_id, rel in d.part.rels.items():
        if "image" in rel.reltype:
            try:
                blob = rel.target_part.blob
            except Exception:
                continue
            ext = rel.target_part.content_type.split("/")[-1]
            doc.images.append(
                ExtractedImage(
                    location="document media",
                    source="content",
                    data=blob,
                    ext=ext,
                    occurrence_hash=hashlib.sha256(blob).hexdigest(),
                )
            )

    doc.unit_count = max(1, len(d.paragraphs) // 40 + 1)  # rough page estimate for volume metrics
    return doc
